from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\1385550\Documents\Codex\Proyectos\el-tarot-de-maria")
OUT = ROOT / "docs" / "auditoria_web_el-tarot-de-maria.dgimdiez77.chatgpt.site_2026-10-13.docx"
NAVY, LIGHT, MID = "17384A", "F4F6F7", "D9D9D9"

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = pr.find(qn("w:shd"))
    if el is None: el = OxmlElement("w:shd"); pr.append(el)
    el.set(qn("w:fill"), fill)

def border(cell):
    pr = cell._tc.get_or_add_tcPr(); edges = pr.first_child_found_in("w:tcBorders")
    if edges is None: edges = OxmlElement("w:tcBorders"); pr.append(edges)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + name); el = edges.find(tag)
        if el is None: el = OxmlElement("w:" + name); edges.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), MID)

def repeat_header(row):
    pr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); pr.append(el)

def keep_row(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])

def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")

def numbered(doc, text):
    doc.add_paragraph(text, style="List Number")

def fill_table(table, headers, rows, widths=None, font=8):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = h; shade(c, NAVY); border(c)
        if widths: c.width = Inches(widths[i])
        for run in c.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255); run.bold = True; run.font.size = Pt(font)
    repeat_header(table.rows[0])
    for idx, data in enumerate(rows):
        row = table.add_row(); keep_row(row)
        for j, value in enumerate(data):
            c = row.cells[j]; c.text = value; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths: c.width = Inches(widths[j])
            shade(c, "FFFFFF" if idx % 2 == 0 else LIGHT); border(c)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs: run.font.size = Pt(font)

doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.62)
sec.left_margin = sec.right_margin = Inches(.72)
normal = doc.styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(10.3); normal.font.color.rgb = RGBColor(32,35,30)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
for name, size in (("Title",27),("Heading 1",18),("Heading 2",13),("Heading 3",11)):
    s = doc.styles[name]; s.font.name = "Aptos Display"; s.font.size = Pt(size); s.font.color.rgb = RGBColor(23,56,74)
    s.font.bold = name != "Title"; s.paragraph_format.keep_with_next = True; s.paragraph_format.space_before = Pt(12 if name != "Title" else 0); s.paragraph_format.space_after = Pt(5)
head = sec.header.paragraphs[0]; head.text = "EL TAROT DE MARÍA  ·  AUDITORÍA DE EXPERIENCIA WEB"; head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
head.runs[0].font.name = "Aptos"; head.runs[0].font.size = Pt(8); head.runs[0].font.color.rgb = RGBColor(91,105,111)
page_number(sec.footer.paragraphs[0])

title = doc.add_paragraph(style="Title"); title.add_run("Auditoría completa de la web\n")
r = title.add_run("EL TAROT DE MARÍA"); r.font.color.rgb = RGBColor(216,168,62)
p = doc.add_paragraph(); r = p.add_run("Versión pública 26 · 13 de octubre de 2026"); r.bold = True; r.font.size = Pt(12)
p = doc.add_paragraph("Enlace público: https://el-tarot-de-maria.dgimdiez77.chatgpt.site/"); p.runs[0].font.color.rgb = RGBColor(37,90,114)

doc.add_heading("Veredicto", level=1)
p = doc.add_paragraph(); p.add_run("El juego ya ofrece una experiencia reconocible, coherente y suficientemente completa para probar retención real con usuarios. ").bold = True
p.add_run("La prioridad inmediata está en hacer más legible la tirada de una carta en móvil, robustecer el diario y explicar mejor cuándo empieza la personalización. La identidad visual, la privacidad local y la facilidad para empezar son fortalezas claras.")

doc.add_heading("Qué se ha revisado", level=2)
for t in [
    "Portada, elección de tema, pregunta opcional, tirada de una carta y tirada de tres cartas.",
    "Revelado, veredicto, guardado, diario, memoria, instalación, privacidad y enlace de WhatsApp.",
    "Vista de escritorio y vista móvil de 390 × 844 píxeles.",
    "Código publicado, metadatos, caché sin conexión y peso de los recursos.",
    "Comportamiento del enlace anterior y del nuevo enlace de Sites.",
]: bullet(doc,t)

doc.add_heading("Fortalezas confirmadas", level=1)
strengths = [
    ("Comprensión inmediata", "El nombre, la pregunta y las dos formas de tirada aparecen sin obligar a registrarse."),
    ("Identidad propia", "Las 23 cartas comparten el rostro de María y una estética uniforme; Géminis queda integrado como carta especial."),
    ("Recorrido completo", "La lectura termina en un veredicto y permite guardar, consultar el diario o compartir el resultado."),
    ("Privacidad clara", "Las preguntas y lecturas permanecen en el navegador. La memoria empieza desactivada y puede borrarse."),
    ("Buen uso móvil", "La tirada de tres cartas usa desplazamiento horizontal, muestra una indicación visible y evita comprimir las ilustraciones."),
    ("Instalable y sin conexión", "El manifiesto, los iconos y la caché preparan el juego como PWA; las pruebas automatizadas de los 31 recursos pasan."),
    ("Límites responsables", "El juego se presenta como lectura simbólica y evita prometer resultados clínicos o certezas futuras."),
]
for a,b in strengths:
    p=doc.add_paragraph(); p.add_run(a+": ").bold=True; p.add_run(b)

doc.add_page_break(); doc.add_heading("Prioridad recomendada", level=1)
pri_rows = [
    ("1","Texto móvil de la tirada de una carta","La explicación baja a 10 px; es el problema de lectura más claro.","Bajo"),
    ("2","Evitar guardados duplicados","Ahora la misma tirada puede añadirse varias veces al diario.","Bajo"),
    ("3","Errores de almacenamiento comprensibles","Si el navegador bloquea o llena el almacenamiento, falta una respuesta visible.","Medio"),
    ("4","Diario preparado para crecer","Filtros, edición y exportación aumentan la utilidad cuando hay muchas lecturas.","Medio"),
    ("5","Explicar el progreso de la memoria","Conviene indicar cuántas lecturas faltan para empezar a personalizar.","Bajo"),
]
fill_table(doc.add_table(rows=1,cols=4),["Orden","Mejora","Por qué ahora","Esfuerzo"],pri_rows,[.55,2.05,3.25,.8],8.7)

doc.add_heading("Hallazgos detallados", level=1)
findings = [
    ("1","Hecho","Bajada al veredicto demasiado rápida","El movimiento automático se percibía como un salto.","Versión 26: recorrido de 1,6 s, foco al terminar y respeto a la reducción de movimiento.","Alto","Bajo"),
    ("2","Alta","Texto pequeño en una carta","En 390 px, el significado de una tirada de una carta se renderiza a 10 px; el veredicto queda a 15 px.","Subir significado y señal a 14 px, con interlineado mínimo de 1,45.","Alto","Bajo"),
    ("3","Alta","Guardado duplicado","El botón vuelve a abrir el diálogo tras guardar y cada confirmación crea otra entrada.","Marcar la lectura actual como guardada, cambiar el texto del botón y permitir actualizar su nota.","Medio","Bajo"),
    ("4","Alta","Fallo silencioso de almacenamiento","Las escrituras en almacenamiento local no capturan errores de cuota o bloqueo.","Mostrar un aviso y mantener la lectura visible para compartirla o copiarla.","Alto","Medio"),
    ("5","Alta","Diario sin herramientas de crecimiento","Con muchas entradas solo se puede recorrer, marcar favorita o borrar.","Añadir filtros por tema y favoritas, búsqueda y orden.","Alto","Medio"),
    ("6","Media","Sin copia ni traslado del diario","Las lecturas quedan ligadas a un navegador y se pierden al borrar sus datos o cambiar de móvil.","Añadir exportación e importación de un archivo JSON, voluntaria y explicada.","Alto","Medio"),
    ("7","Media","Inicio de la memoria poco concreto","La interfaz no dice claramente que necesita dos lecturas guardadas.","Mostrar “1 de 2 lecturas” y después resumir los temas detectados.","Medio","Bajo"),
    ("8","Media","Instalación genérica","El diálogo remite al menú del navegador y no explica primero la ventaja de instalar.","Detectar Android/iPhone, explicar acceso rápido y modo sin conexión, y dar pasos breves.","Medio","Bajo"),
    ("9","Media","Sin medición del recorrido","No se sabe cuántas personas empiezan, terminan, guardan, vuelven o comparten.","Definir cinco eventos mínimos y elegir medición respetuosa con la privacidad.","Alto","Medio"),
    ("10","Media","Caché inicial pesada","El sitio ocupa 11,8 MB y las 23 cartas suman 7,7 MB; se preparan todas al instalarse.","Cargar portada y núcleo primero; guardar el resto después o a demanda.","Medio","Medio"),
    ("11","Media","SEO técnico incompleto","Hay título, descripción y Open Graph, pero no canonical, robots.txt, sitemap.xml ni datos estructurados.","Añadir canonical, robots, sitemap y esquema WebApplication o Game.","Medio","Bajo"),
    ("12","Baja","Botones de borrado poco específicos","Cada entrada usa el nombre accesible “Borrar”; no identifica la lectura.","Usar “Borrar lectura del 12 de octubre” mediante aria-label.","Medio","Bajo"),
    ("13","Baja","Atrás no representa el estado del juego","Inicio, tirada y diálogos no se registran en el historial del navegador.","Hacer que Atrás vuelva a la portada o cierre el diálogo antes de salir.","Medio","Medio"),
]
fill_table(doc.add_table(rows=1,cols=7),["#","Nivel","Hallazgo","Evidencia","Corrección","Impacto","Esfuerzo"],findings,[.28,.62,1.05,1.55,2.0,.55,.55],7.2)

doc.add_heading("Cinco mejoras rápidas", level=1)
for t in [
    "Subir a 14 px el texto del significado en la tirada móvil de una carta.",
    "Desactivar “Guardar en mi diario” después de guardar y ofrecer “Editar nota”.",
    "Mostrar el progreso de la memoria: 0 de 2, 1 de 2 y activa.",
    "Hacer específicos los nombres accesibles de Borrar y Favorita.",
    "Añadir canonical, robots.txt y sitemap.xml con la dirección nueva.",
]: numbered(doc,t)

doc.add_heading("Hoja de ruta sugerida", level=1)
for a,b in [
    ("Tanda 1 · Claridad y fiabilidad","Texto móvil, guardado único, progreso de memoria, aviso de error al guardar e instrucciones por dispositivo."),
    ("Tanda 2 · Diario útil","Filtros, búsqueda, edición de notas, exportación e importación y estados de historial."),
    ("Tanda 3 · Crecimiento medible","Carga progresiva de imágenes, SEO técnico, prueba con usuarios y medición mínima con consentimiento."),
]:
    p=doc.add_paragraph(); p.add_run(a+"\n").bold=True; p.add_run(b)

doc.add_heading("Cambio de enlace y dominios", level=1)
p=doc.add_paragraph(); p.add_run("Nuevo enlace de Sites: ").bold=True; p.add_run("https://el-tarot-de-maria.dgimdiez77.chatgpt.site/")
p=doc.add_paragraph(); p.add_run("Enlace anterior: ").bold=True; p.add_run("responde con una redirección temporal 307 al nuevo enlace. Los enlaces ya compartidos continúan llegando al juego.")
domains = [
    ("eltarotdemaria.com","No disponible","Registro oficial .com: alta el 03/12/2002 y vencimiento registrado el 03/12/2026.","Requeriría negociación con el titular."),
    ("eltarotdemaria.es","No disponible","Tiene servidores DNS autoritativos de canaldominios.com; el nombre está delegado.","No intentar comprarlo como dominio libre."),
    ("eltarotdemaría.com","Parece disponible","La forma IDN devuelve “no encontrado” en RDAP .com y no presenta servidores DNS.","Confirmar en el registrador. Forma técnica: xn--eltarotdemara-7ib.com."),
    ("eltarotdemaría.es","Parece disponible","No presenta servidores DNS en la consulta realizada.","Confirmar en Dominios.es. Forma técnica: xn--eltarotdemara-7ib.es."),
]
fill_table(doc.add_table(rows=1,cols=4),["Nombre consultado","Resultado","Evidencia","Recomendación"],domains,[1.25,1.0,2.7,2.15],8)
p=doc.add_paragraph(); p.add_run("Criterio: ").bold=True; p.add_run("conviene usar un dominio sin tilde como dirección principal porque es más fácil de dictar, escribir y compartir. Como los dos nombres sin tilde están ocupados, antes de comprar la versión acentuada merece la pena comparar alternativas cortas sin tilde.")

doc.add_heading("Fuentes y límites", level=1)
for t in [
    "Revisión directa del sitio público y de sus pantallas el 13 de octubre de 2026.",
    "Código correspondiente a la versión 26 y cuatro pruebas automatizadas superadas.",
    "Verisign RDAP para .com: https://rdap.verisign.com/com/v1/domain/eltarotdemaria.com",
    "Dominios.es como buscador oficial de .es: https://www.dominios.es/es",
    "Consultas DNS para .es y para las variantes con tilde convertidas a IDN.",
]: bullet(doc,t)
p=doc.add_paragraph(); p.add_run("Límite: ").bold=True; p.add_run("la disponibilidad de un dominio puede cambiar en cualquier momento. Una consulta no lo reserva. El precio y la disponibilidad deben confirmarse en el registrador elegido antes de comprar.")

doc.core_properties.title="Auditoría web de EL TAROT DE MARÍA"; doc.core_properties.subject="Experiencia de usuario, accesibilidad, conversión, SEO y dominios"; doc.core_properties.author="Codex"
doc.save(OUT); print(OUT)
